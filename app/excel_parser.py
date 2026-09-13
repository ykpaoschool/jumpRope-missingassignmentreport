"""解析 JumpRope 导出的 Excel，按学生分组。"""

from __future__ import annotations

import datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

# 表头名 -> 内部字段名
COL_HEADERS = {
    "student_id": "Student External Id",
    "grade": "Student Current Grade Level",
    "class": "Student Official Class",
    "school": "School Short Code",
    "course": "Section Course Name",
    "teacher": "Section Teacher Name",
    "title": "Assessment Title",
    "due_date": "Assessment Due Date",
    "code": "Missing Work Code",
}

# 用于定位家长邮箱列的表头关键词
EMAIL_KEYWORDS = ("email", "mail", "邮箱", "家长")

# 必需列缺失时视为无法解析
REQUIRED_HEADERS = ("student_id", "course", "teacher", "title", "due_date", "code")


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_date_str(value) -> str:
    """把截止日期统一为 'YYYY-MM-DD' 字符串。"""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        # Excel 日期序列号（如 46273 -> 2026-09-08）
        try:
            return from_excel(value).strftime("%Y-%m-%d")
        except Exception:  # noqa: BLE001
            return _clean(value)
    s = _clean(value)
    if not s:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%Y.%m.%d"):
        try:
            return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s


def parse_workbook(data: bytes, filename: str = "") -> dict:
    """解析 Excel 字节，返回 {'students': [...], 'warnings': [...]}。"""
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"students": [], "warnings": ["文件为空，无数据"]}

    header = [_clean(h) for h in rows[0]]

    idx: dict[str, int | None] = {}
    for key, name in COL_HEADERS.items():
        idx[key] = header.index(name) if name in header else None

    # 家长邮箱列：优先按表头关键词定位，否则回退到第 14 列（N）
    email_col: int | None = None
    for i, h in enumerate(header):
        low = h.lower()
        if any(k in low for k in EMAIL_KEYWORDS):
            email_col = i
            break
    if email_col is None and len(header) >= 14:
        email_col = 13

    missing = [name for key, name in COL_HEADERS.items() if idx[key] is None and key in REQUIRED_HEADERS]
    if missing:
        return {"students": [], "warnings": [f"缺少必需列：{'、'.join(missing)}"]}

    def get(row, key):
        i = idx.get(key)
        return row[i] if i is not None and i < len(row) else None

    grouped: dict[str, dict] = {}
    order: list[str] = []
    warnings: list[str] = []

    for row in rows[1:]:
        if row is None or all(v is None or _clean(v) == "" for v in row):
            continue
        sid = _clean(get(row, "student_id"))
        if not sid:
            continue
        email = _clean(row[email_col]) if email_col is not None and email_col < len(row) else ""

        if sid not in grouped:
            grouped[sid] = {
                "student_id": sid,
                "grade": _clean(get(row, "grade")),
                "class": _clean(get(row, "class")),
                "school": _clean(get(row, "school")),
                "parent_email": email,
                "items": [],
            }
            order.append(sid)
        else:
            stu = grouped[sid]
            if email and stu["parent_email"] and email.lower() != stu["parent_email"].lower():
                warnings.append(f"学生 {sid} 出现多个不同家长邮箱：{stu['parent_email']} / {email}")
            if email and not stu["parent_email"]:
                stu["parent_email"] = email

        grouped[sid]["items"].append(
            {
                "course": _clean(get(row, "course")),
                "teacher": _clean(get(row, "teacher")),
                "title": _clean(get(row, "title")),
                "due_date": _to_date_str(get(row, "due_date")),
                "code": _clean(get(row, "code")),
            }
        )

    students = [grouped[sid] for sid in order]
    no_email = [s["student_id"] for s in students if not s["parent_email"]]
    if no_email:
        warnings.append(f"以下 {len(no_email)} 名学生缺少家长邮箱，将不会被发送：{', '.join(no_email)}")

    return {"students": students, "warnings": warnings}
